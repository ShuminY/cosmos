"""PDF / Word / PowerPoint analyzers — extract text + structure."""
from __future__ import annotations
from pathlib import Path

from pypdf import PdfReader
import docx
import pptx


MAX_TEXT_CHARS = 50_000


def analyze_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    n_pages = len(reader.pages)
    chunks = []
    for i, page in enumerate(reader.pages):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        chunks.append(t)
        if sum(len(c) for c in chunks) >= MAX_TEXT_CHARS:
            break
    text = "\n\n".join(chunks)[:MAX_TEXT_CHARS]
    meta = reader.metadata or {}
    return {
        "format": "pdf",
        "summary": f"{n_pages} page(s), {len(text):,} chars extracted",
        "n_pages": n_pages,
        "title": str(meta.get("/Title", "")) if meta else "",
        "author": str(meta.get("/Author", "")) if meta else "",
        "text_preview": text[:5000],
        "text": text,   # full extracted (truncated to MAX_TEXT_CHARS)
    }


def analyze_docx(path: Path) -> dict:
    d = docx.Document(str(path))
    paragraphs = [p.text for p in d.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)[:MAX_TEXT_CHARS]
    n_tables = len(d.tables)
    return {
        "format": "docx",
        "summary": f"{len(paragraphs)} paragraph(s), {n_tables} table(s), {len(text):,} chars",
        "n_paragraphs": len(paragraphs),
        "n_tables": n_tables,
        "text_preview": text[:5000],
        "text": text,
    }


def analyze_pptx(path: Path) -> dict:
    prs = pptx.Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides):
        title = ""
        body_chunks = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if not txt:
                    continue
                if not title and shape == slide.shapes.title:
                    title = txt
                else:
                    body_chunks.append(txt)
        slides.append({
            "index": i,
            "title": title,
            "body": "\n".join(body_chunks)[:2000],
        })
    text = "\n\n".join(f"[Slide {s['index']+1}] {s['title']}\n{s['body']}" for s in slides)[:MAX_TEXT_CHARS]
    return {
        "format": "pptx",
        "summary": f"{len(slides)} slide(s), {len(text):,} chars",
        "n_slides": len(slides),
        "slides": slides[:50],
        "text_preview": text[:5000],
    }
