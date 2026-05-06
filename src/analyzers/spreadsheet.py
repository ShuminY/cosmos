"""Excel / CSV analyzer — list sheets, columns, and small previews."""
from __future__ import annotations
from pathlib import Path

import openpyxl


def analyze_xlsx(path: Path) -> dict:
    if path.suffix.lower() == ".csv":
        # use openpyxl read mode? csv needs different. Use pandas.
        import pandas as pd
        df = pd.read_csv(path, nrows=200)
        return {
            "format": "csv",
            "summary": f"CSV — {len(df)} rows × {len(df.columns)} cols (first 200)",
            "sheets": [{
                "name": "csv",
                "n_rows": len(df),
                "n_cols": len(df.columns),
                "columns": [str(c) for c in df.columns],
                "preview": df.head(20).fillna("").astype(str).to_dict(orient="records"),
            }],
        }

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    sheets = []
    total_rows = 0
    for name in wb.sheetnames:
        ws = wb[name]
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        total_rows += max_row
        # Sample first 30 rows for preview
        rows_iter = ws.iter_rows(values_only=True, max_row=30)
        rows = []
        header = None
        for i, row in enumerate(rows_iter):
            cells = [str(c) if c is not None else "" for c in row]
            if i == 0:
                header = cells
            rows.append(cells)
        sheets.append({
            "name": name,
            "n_rows": max_row,
            "n_cols": max_col,
            "columns": header or [],
            "preview": rows[:20],
        })
    wb.close()
    return {
        "format": "xlsx",
        "summary": f"{len(sheets)} sheet(s), ~{total_rows} total rows",
        "n_sheets": len(sheets),
        "sheets": sheets,
    }
